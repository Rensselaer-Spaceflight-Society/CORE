"""Read-only dump of the October 2 budget workbook (formulas and cached values). Never saves."""
import hashlib, sys
from pathlib import Path
import openpyxl
WB = Path('C:/Users/andyc/OneDrive/Desktop/CORE/outputs/core-budget-20261002-fbef9859/CORE_Itemized_Budget.xlsx')
h0 = hashlib.sha256(WB.read_bytes()).hexdigest()
wf = openpyxl.load_workbook(WB, read_only=True, data_only=False)
wv = openpyxl.load_workbook(WB, read_only=True, data_only=True)
print('sheets', wf.sheetnames)
for name in wf.sheetnames:
    f, v = wf[name], wv[name]
    print('=====', name, f.max_row, f.max_column)
    vrows = list(v.iter_rows(values_only=True))
    for i, row in enumerate(f.iter_rows(values_only=True), start=1):
        vals = vrows[i-1]
        cells = []
        for j, (cf, cv) in enumerate(zip(row, vals), start=1):
            if cf is None and cv is None: continue
            col = openpyxl.utils.get_column_letter(j)
            if isinstance(cf, str) and cf.startswith('='):
                cells.append(f'{col}{i}: {cf} -> {cv!r}')
            else:
                cells.append(f'{col}{i}: {cv!r}')
        if cells and (i <= 8 or i >= 118 or name != wf.sheetnames[0]):
            print(' | '.join(cells))
        elif cells and i in (9, 10, 60):
            print(' | '.join(cells))
wf.close(); wv.close()
h1 = hashlib.sha256(WB.read_bytes()).hexdigest()
print('sha256 before', h0); print('sha256 after ', h1); print('unchanged' if h0 == h1 else 'CHANGED!')
